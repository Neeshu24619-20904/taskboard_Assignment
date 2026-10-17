from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

prs = Presentation()

def add_slide(title_text, bullets):
    slide_layout = prs.slide_layouts[1] # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    
    # Title formatting
    title = slide.shapes.title
    title.text = title_text
    title.text_frame.paragraphs[0].font.size = Pt(28)
    title.text_frame.paragraphs[0].font.bold = True
    title.text_frame.paragraphs[0].font.color.rgb = RGBColor(31, 78, 121)
    
    # Content formatting
    tf = slide.placeholders[1].text_frame
    tf.clear()
    for i, bullet in enumerate(bullets):
        p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
        p.text = bullet
        p.level = 0
        p.font.size = Pt(18)
        p.space_after = Pt(10)

# Slide 1: Title
slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(slide_layout)
slide.shapes.title.text = "TaskBoard — DevOps Capstone Project"
slide.placeholders[1].text = "Name: Neeshant Ahire\nRoll No: 10635\nBatch: A\nEnd-to-End CI/CD, Containerization, Kubernetes & Observability"

# Slide 2
add_slide("Project Overview & Objectives", [
    "Application: SaaS project management tool tracking tasks across workflows.",
    "Core Goal: Demonstrate an automated DevOps lifecycle from code commit to Kubernetes.",
    "Tech Stack: React + Vite, FastAPI, PostgreSQL, Docker, Kubernetes, Helm, Prometheus."
])

# Slide 3
add_slide("End-to-End DevOps Architecture", [
    "Code & Version Control: Git & GitHub repository tracking desired states.",
    "CI/CD Automation: GitHub Actions running automated Pytest & React builds.",
    "Container Registry: Secure image publishing via GitHub Container Registry (GHCR).",
    "Orchestration & Monitoring: Deployed via Helm onto Kubernetes, tracked with Prometheus & Grafana."
])

# Slide 4
add_slide("Application Layer & Quality Gates", [
    "FastAPI Backend: Exposes high-performance REST APIs (/api/tasks, /health, /ready, /metrics).",
    "Database Management: SQLAlchemy ORM backed by PostgreSQL with Alembic migrations.",
    "Automated Testing: Pytest test suite acts as the first strict quality gate in CI."
])

# Slide 5
add_slide("Containerization & Docker", [
    "Backend Dockerfile: Optimized Python runtime container.",
    "Frontend Multi-Stage Build: Node build phase packaged into a lightweight Nginx runtime image.",
    "Security Best Practice: Both application containers run as non-root users.",
    "Local Execution: Single-command deployment via Docker Compose."
])

# Slide 6
add_slide("CI/CD Pipeline & DevSecOps", [
    "GitHub Actions Pipeline: Automated execution on pushes to the main branch.",
    "Immutable Tagging: Images are tagged with the Git commit SHA instead of 'latest'.",
    "Trivy Security Scanning: Automatic scan for HIGH and CRITICAL container vulnerabilities."
])

# Slide 7
add_slide("Infrastructure as Code (Terraform) & AWS", [
    "Declarative Infrastructure: Provisioning AWS resources using Terraform HCL scripts.",
    "Cloud Architecture: Configured for AWS VPC with subnets and an Amazon EKS cluster.",
    "Security Compliance: Credentials secured via variables, avoiding raw keys in git history."
])

# Slide 8
add_slide("Kubernetes & Helm Orchestration", [
    "Namespaces & Deployments: Isolated workloads managed securely inside Kubernetes.",
    "Helm Packaging: Unified configuration charts replacing manual multi-file manifests.",
    "Ingress Controllers: Efficient path-based routing for frontend and backend traffic."
])

# Slide 9
add_slide("Observability & GitOps", [
    "Prometheus Monitoring: Automatic scraping of FastAPI metrics via /metrics and ServiceMonitors.",
    "Grafana Dashboards: Live visual metrics tracking resource usage, latencies, and requests.",
    "GitOps Alignment: Version-controlled desired states ensuring reproducibility."
])

# Slide 10
add_slide("Conclusion & Deliverables", [
    "Successfully fulfilled all 10 grading modules (M1-M10).",
    "Fully traceable pipeline from local workstation to production Kubernetes cluster.",
    "GitHub Repository, Documentation, and Demo Presentation fully verified."
])

prs.save("TaskBoard_Capstone_Presentation.pptx")
print("Presentation generated successfully as 'TaskBoard_Capstone_Presentation.pptx'!")
